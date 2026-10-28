"""
Shared helpers for writing readable Excel inspection copies.

These workbooks are HUMAN-READABLE COPIES ONLY. The authoritative data stays in
analysis_data/*.parquet; the ML pipeline must keep reading Parquet, because
Excel is far too slow for hundreds of thousands of rows.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXCEL_MAX_ROWS = 1_048_576
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
TS_FMT = "yyyy-mm-dd hh:mm:ss"
FLOAT_DECIMALS = 6          # kills float noise (1.7000000000000002 -> 1.7)


def _width(col, series):
    """Column width from the header and a sample of the values."""
    n = len(str(col))
    if pd.api.types.is_datetime64_any_dtype(series):
        return 20
    sample = series.head(400)
    if len(sample):
        try:
            n = max(n, int(sample.astype(str).str.len().max()))
        except Exception:
            pass
    return max(9, min(n + 3, 34))


def write_table_sheet(wb, df, title, freeze="A2"):
    """
    Append one sheet holding `df`: styled header, frozen header row, autofilter,
    readable widths, real datetimes formatted as yyyy-mm-dd hh:mm:ss.

    Uses openpyxl write-only mode - the only mode fast enough at this row count.
    """
    ws = wb.create_sheet(title=title[:31])
    ws.freeze_panes = freeze

    for i, c in enumerate(df.columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _width(c, df[c])

    header = []
    for c in df.columns:
        cell = WriteOnlyCell(ws, value=str(c))
        cell.font, cell.fill = HDR_FONT, HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.append(cell)
    ws.append(header)

    dt_cols = {c for c in df.columns if pd.api.types.is_datetime64_any_dtype(df[c])}
    # Round floats once, vectorised, instead of styling millions of cells.
    out = df.copy()
    for c in out.columns:
        if pd.api.types.is_float_dtype(out[c]):
            out[c] = out[c].round(FLOAT_DECIMALS)
    out = out.astype(object).where(pd.notna(out), None)

    cols = list(out.columns)
    for row in out.itertuples(index=False, name=None):
        cells = []
        for c, v in zip(cols, row):
            if c in dt_cols and v is not None:
                cell = WriteOnlyCell(ws, value=v)
                cell.number_format = TS_FMT
                cells.append(cell)
            else:
                cells.append(v)
        ws.append(cells)

    ws.auto_filter.ref = "A1:{}{}".format(get_column_letter(len(cols)), len(out) + 1)
    return ws


def write_kv_sheet(wb, rows, title, col_widths=(46, 62), header=None):
    """
    A simple 'Field | Value' summary sheet. Works in both write-only and normal
    workbooks. Rows whose first cell starts with '--' are rendered as bold
    section headings.
    """
    ws = wb.create_sheet(title=title[:31])
    ws.freeze_panes = "A2"
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    h = list(header or ("Field", "Value"))
    hdr = []
    for v in h:
        cell = WriteOnlyCell(ws, value=v) if wb.write_only else None
        if cell is None:
            hdr.append(v)
        else:
            cell.font, cell.fill = HDR_FONT, HDR_FILL
            cell.alignment = Alignment(horizontal="center")
            hdr.append(cell)
    ws.append(hdr)
    if not wb.write_only:
        for c in ws[1]:
            c.font, c.fill = HDR_FONT, HDR_FILL
            c.alignment = Alignment(horizontal="center")

    for r in rows:
        r = list(r)
        is_head = r and isinstance(r[0], str) and r[0].startswith("--")
        if wb.write_only:
            out = []
            for v in r:
                cell = WriteOnlyCell(ws, value=v)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if is_head:
                    cell.font = Font(bold=True)
                out.append(cell)
            ws.append(out)
        else:
            ws.append(r)
    if not wb.write_only:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
            if row[0].value and str(row[0].value).startswith("--"):
                for c in row:
                    c.font = Font(bold=True)
    return ws


def new_workbook(write_only=True):
    wb = Workbook(write_only=write_only)
    return wb


def finalize(wb, path):
    """Drop the default empty sheet (normal mode only) and save."""
    if not wb.write_only and "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    wb.save(path)
    return path
